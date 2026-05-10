#!/usr/bin/env python3
"""Loop tests against the packaged picvert-engine binary.

This is a black-box test of the *built* sidecar binary — i.e. exactly what the
Tauri shell calls. Catches packaging issues that source-level pytest misses
(missing PyInstaller hidden imports, runtime path mistakes, etc.).

Run:
    python scripts/loop_test_engine.py --rounds 5
    python scripts/loop_test_engine.py --rounds 5 --binary src-tauri/binaries/picvert-engine-aarch64-apple-darwin

Each round runs every test case once. Failures are collected and summarised at
the end. Exit code 0 iff every round was clean.
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import subprocess
import sys
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook


# ─────────────────────────────────────────────────── infrastructure

@dataclass
class CaseResult:
    name: str
    ok: bool
    detail: str = ""
    duration_ms: float = 0.0


@dataclass
class RoundResult:
    round_num: int
    cases: list[CaseResult] = field(default_factory=list)

    @property
    def passed(self) -> int:
        return sum(1 for c in self.cases if c.ok)

    @property
    def failed(self) -> int:
        return sum(1 for c in self.cases if not c.ok)


def call_engine(binary: Path, request: dict) -> dict:
    """Invoke the engine in --json one-shot mode, return parsed reply."""
    payload = json.dumps(request)
    # PyInstaller --onefile cold start is ~6-7s on macOS (extracts to /tmp).
    # Tests should be lenient; the real fix is a persistent sidecar in Rust.
    proc = subprocess.run(
        [str(binary), "--json", payload],
        capture_output=True,
        text=True,
        timeout=60,
    )
    if proc.returncode not in (0, 2):
        raise RuntimeError(
            f"engine exited with status {proc.returncode}\n"
            f"stderr: {proc.stderr.strip()}\n"
            f"stdout: {proc.stdout.strip()}"
        )
    last_line = [l for l in proc.stdout.strip().splitlines() if l][-1]
    return json.loads(last_line)


# ─────────────────────────────────────────────────── fixtures

def make_png(path: Path, color=(255, 0, 0), size=(40, 40)) -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def make_csv(path: Path, rows=(("a", "b"), ("1", "2"))) -> Path:
    with path.open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)
    return path


def make_xlsx(path: Path, rows=(("h1", "h2"), ("v1", "v2")), sheet="First") -> Path:
    wb = Workbook()
    wb.active.title = sheet
    for r in rows:
        wb.active.append(r)
    wb.save(str(path))
    return path


def make_docx(path: Path, paras=("hello world", "second line")) -> Path:
    doc = Document()
    for p in paras:
        doc.add_paragraph(p)
    doc.save(str(path))
    return path


def make_pdf_2page(path: Path) -> Path:
    """Build a 2-page PDF using PyMuPDF (Pillow's multi-page PDF writer trips
    on the JPEG codec; fitz is what the engine uses anyway, so this matches
    real input shape too)."""
    import fitz
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page(width=72, height=72)
        page.draw_rect(fitz.Rect(10, 10, 50, 50), fill=(0.2, 0.6, 1.0))
    doc.save(str(path))
    doc.close()
    return path


# ─────────────────────────────────────────────────── test cases

def case_ping(binary: Path) -> tuple[bool, str]:
    r = call_engine(binary, {"id": "p1", "action": "ping"})
    if not (r.get("ok") and r.get("result", {}).get("pong") is True):
        return False, f"unexpected reply: {r}"
    return True, ""


def case_list_formats(binary: Path) -> tuple[bool, str]:
    r = call_engine(binary, {"id": "f1", "action": "list_formats"})
    if not r.get("ok"):
        return False, f"failed: {r}"
    res = r["result"]
    for ext in (".png", ".pdf", ".xlsx", ".docx", ".csv", ".heic"):
        if ext not in res["input_extensions"]:
            return False, f"missing input ext {ext}"
    for fmt in ("PNG", "JPEG", "PDF", "DOCX", "XLSX", "CSV", "HEIC", "SVG"):
        if fmt not in res["output_formats"]:
            return False, f"missing output format {fmt}"
    return True, ""


def case_unknown_action(binary: Path) -> tuple[bool, str]:
    r = call_engine(binary, {"action": "fly_to_mars"})
    if r.get("ok") is not False:
        return False, "expected ok=false"
    if "unknown_action" not in r["error"]["message"]:
        return False, f"wrong error: {r['error']}"
    return True, ""


def case_unknown_format(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_png(tmp / "in.png")
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(tmp / "out"),
        "format": "BOGUS",
    })
    if r.get("ok") is not False:
        return False, f"expected ok=false, got {r}"
    if r["error"]["kind"] != "unsupported":
        return False, f"wrong kind: {r['error']}"
    return True, ""


def _make_image_conversion_case(out_format: str, ext: str) -> Callable:
    def fn(binary: Path, tmp: Path) -> tuple[bool, str]:
        src = make_png(tmp / "in.png")
        out = tmp / "out"
        r = call_engine(binary, {
            "action": "convert", "input": str(src), "output": str(out),
            "format": out_format,
        })
        if not r.get("ok"):
            return False, f"engine error: {r}"
        target = out / f"in{ext}"
        if not target.exists():
            return False, f"output file missing: {target}"
        return True, ""
    fn.__name__ = f"case_png_to_{out_format.lower()}"
    return fn


def case_pdf_to_png(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_pdf_2page(tmp / "doc.pdf")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "PNG",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    if r["result"]["written"] != 2:
        return False, f"expected 2 pages written, got {r['result']['written']}"
    if not (out / "doc_page1.png").exists():
        return False, "page 1 missing"
    if not (out / "doc_page2.png").exists():
        return False, "page 2 missing"
    return True, ""


def case_csv_to_xlsx(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_csv(tmp / "in.csv")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "XLSX",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    target = out / "in.xlsx"
    if not target.exists():
        return False, "output xlsx missing"
    wb = load_workbook(target)
    rows = [tuple(row) for row in wb.active.iter_rows(values_only=True)]
    if rows != [("a", "b"), ("1", "2")]:
        return False, f"unexpected rows: {rows}"
    return True, ""


def case_csv_to_docx(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_csv(tmp / "in.csv")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "DOCX",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    doc = Document(str(out / "in.docx"))
    if len(doc.tables) != 1:
        return False, f"expected 1 table, got {len(doc.tables)}"
    return True, ""


def case_xlsx_to_csv_single(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_xlsx(tmp / "in.xlsx")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "CSV",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    if r["result"]["written"] != 1:
        return False, f"expected 1 written, got {r['result']['written']}"
    return True, ""


def case_xlsx_to_csv_multi(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = tmp / "multi.xlsx"
    wb = Workbook()
    wb.active.title = "Alpha"
    wb.active.append(["a", "1"])
    wb.create_sheet("Beta").append(["b", "2"])
    wb.create_sheet("Gamma").append(["c", "3"])
    wb.save(str(src))
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "CSV",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    if r["result"]["written"] != 3:
        return False, f"expected 3 csv files, got {r['result']['written']}"
    for name in ("multi__Alpha.csv", "multi__Beta.csv", "multi__Gamma.csv"):
        if not (out / name).exists():
            return False, f"missing {name}"
    return True, ""


def case_docx_to_xlsx(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_docx(tmp / "in.docx")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "XLSX",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    wb = load_workbook(out / "in.xlsx")
    if "Paragraphs" not in wb.sheetnames:
        return False, f"expected Paragraphs sheet, got {wb.sheetnames}"
    return True, ""


def case_docx_to_csv(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_docx(tmp / "in.docx")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "CSV",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    return True, ""


def case_same_format_copy(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_csv(tmp / "in.csv")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "CSV",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    if (out / "in.csv").read_text() != src.read_text():
        return False, "content mismatch after copy"
    return True, ""


def case_truncated_png(binary: Path, tmp: Path) -> tuple[bool, str]:
    bad = tmp / "bad.png"
    bad.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    r = call_engine(binary, {
        "action": "convert", "input": str(bad), "output": str(tmp / "out"),
        "format": "JPEG",
    })
    if r.get("ok") is not False:
        return False, f"expected ok=false for truncated png, got {r}"
    return True, ""


def case_missing_input(binary: Path, tmp: Path) -> tuple[bool, str]:
    r = call_engine(binary, {
        "action": "convert", "input": "/this/does/not/exist.png",
        "output": str(tmp / "out"), "format": "JPEG",
    })
    if r.get("ok") is not False:
        return False, f"expected ok=false for missing file, got {r}"
    return True, ""


def case_invalid_json(binary: Path) -> tuple[bool, str]:
    proc = subprocess.run(
        [str(binary), "--json", "{not json}"],
        capture_output=True, text=True, timeout=60,
    )
    if proc.returncode != 2:
        return False, f"expected exit 2, got {proc.returncode}"
    last = [l for l in proc.stdout.strip().splitlines() if l][-1]
    resp = json.loads(last)
    if resp.get("ok") is not False:
        return False, f"expected ok=false, got {resp}"
    if resp["error"]["kind"] != "bad_request":
        return False, f"expected bad_request, got {resp['error']}"
    return True, ""


def case_line_mode(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_png(tmp / "in.png")
    requests = "\n".join([
        '{"id":"a","action":"ping"}',
        json.dumps({
            "id": "b", "action": "convert",
            "input": str(src), "output": str(tmp / "out"), "format": "PNG",
        }),
        '{"id":"c","action":"shutdown"}',
    ]) + "\n"
    proc = subprocess.run(
        [str(binary), "--line-mode"],
        input=requests, capture_output=True, text=True, timeout=60,
    )
    if proc.returncode != 0:
        return False, f"exit {proc.returncode}: {proc.stderr}"
    lines = [json.loads(l) for l in proc.stdout.strip().splitlines() if l]
    if len(lines) != 3:
        return False, f"expected 3 replies, got {len(lines)}"
    if lines[0]["id"] != "a" or not lines[0]["ok"]:
        return False, f"ping bad: {lines[0]}"
    if lines[1]["id"] != "b" or lines[1]["result"]["written"] != 1:
        return False, f"convert bad: {lines[1]}"
    if lines[2]["id"] != "c":
        return False, f"shutdown bad: {lines[2]}"
    return True, ""


def case_heic_round_trip(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_png(tmp / "src.png", color=(10, 200, 30))
    r1 = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(tmp), "format": "HEIC",
    })
    if not r1.get("ok"):
        return False, f"png→heic failed: {r1}"
    heic = tmp / "src.heic"
    if not heic.exists():
        return False, "heic missing"
    r2 = call_engine(binary, {
        "action": "convert", "input": str(heic), "output": str(tmp / "back"),
        "format": "PNG",
    })
    if not r2.get("ok"):
        return False, f"heic→png failed: {r2}"
    if not (tmp / "back" / "src.png").exists():
        return False, "round-trip png missing"
    return True, ""


def case_svg_output(binary: Path, tmp: Path) -> tuple[bool, str]:
    src = make_png(tmp / "in.png")
    out = tmp / "out"
    r = call_engine(binary, {
        "action": "convert", "input": str(src), "output": str(out), "format": "SVG",
    })
    if not r.get("ok"):
        return False, f"engine error: {r}"
    text = (out / "in.svg").read_text()
    if not text.startswith("<?xml") or "image/png;base64" not in text:
        return False, "svg looks malformed"
    return True, ""


# ─────────────────────────────────────────────────── runner

def all_cases() -> list[tuple[str, Callable]]:
    cases: list[tuple[str, Callable]] = [
        ("ping", lambda b, _t: case_ping(b)),
        ("list_formats", lambda b, _t: case_list_formats(b)),
        ("unknown_action", lambda b, _t: case_unknown_action(b)),
        ("invalid_json", lambda b, _t: case_invalid_json(b)),
        ("unknown_format", case_unknown_format),
        ("missing_input", case_missing_input),
        ("truncated_png", case_truncated_png),
        ("png→jpeg", _make_image_conversion_case("JPEG", ".jpg")),
        ("png→webp", _make_image_conversion_case("WEBP", ".webp")),
        ("png→tiff", _make_image_conversion_case("TIFF", ".tiff")),
        ("png→bmp", _make_image_conversion_case("BMP", ".bmp")),
        ("png→gif", _make_image_conversion_case("GIF", ".gif")),
        ("png→ico", _make_image_conversion_case("ICO", ".ico")),
        ("png→pdf", _make_image_conversion_case("PDF", ".pdf")),
        ("png→svg", case_svg_output),
        ("heic_round_trip", case_heic_round_trip),
        ("pdf→png", case_pdf_to_png),
        ("csv→xlsx", case_csv_to_xlsx),
        ("csv→docx", case_csv_to_docx),
        ("xlsx→csv (single sheet)", case_xlsx_to_csv_single),
        ("xlsx→csv (multi sheet)", case_xlsx_to_csv_multi),
        ("docx→xlsx", case_docx_to_xlsx),
        ("docx→csv", case_docx_to_csv),
        ("csv→csv (same format copy)", case_same_format_copy),
        ("line-mode round trip", case_line_mode),
    ]
    return cases


def run_round(round_num: int, binary: Path) -> RoundResult:
    res = RoundResult(round_num=round_num)
    cases = all_cases()
    for name, fn in cases:
        with tempfile.TemporaryDirectory(prefix=f"picvert_r{round_num}_") as td:
            tmp = Path(td)
            import time
            t0 = time.perf_counter()
            try:
                ok, detail = fn(binary, tmp)
            except Exception:
                ok, detail = False, traceback.format_exc()
            t1 = time.perf_counter()
            res.cases.append(CaseResult(
                name=name, ok=ok, detail=detail, duration_ms=(t1 - t0) * 1000,
            ))
    return res


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rounds", type=int, default=3, help="how many full passes (default: 3)")
    parser.add_argument(
        "--binary",
        type=Path,
        default=Path("dist/picvert-engine"),
        help="path to the picvert-engine binary",
    )
    args = parser.parse_args()

    if not args.binary.exists():
        print(f"❌ binary not found: {args.binary}", file=sys.stderr)
        return 1

    print(f"▶ Loop-testing {args.binary}  ×  {args.rounds} round(s)")
    rounds = [run_round(i + 1, args.binary) for i in range(args.rounds)]

    total_passed = sum(r.passed for r in rounds)
    total_failed = sum(r.failed for r in rounds)
    total_cases = sum(len(r.cases) for r in rounds)

    # Per-case breakdown across rounds.
    print("\n┌" + "─" * 60 + "┬" + "─" * 12 + "┐")
    print(f"│ {'case':<58} │ {'pass/total':>10} │")
    print("├" + "─" * 60 + "┼" + "─" * 12 + "┤")
    case_stats: dict[str, list[bool]] = {}
    for r in rounds:
        for c in r.cases:
            case_stats.setdefault(c.name, []).append(c.ok)
    for name, results in case_stats.items():
        passed = sum(1 for r in results if r)
        total = len(results)
        marker = "✅" if passed == total else "❌"
        print(f"│ {marker} {name:<56} │ {passed:>3}/{total:<6} │")
    print("└" + "─" * 60 + "┴" + "─" * 12 + "┘")

    if total_failed > 0:
        print("\n=== FAILURES ===")
        for r in rounds:
            for c in r.cases:
                if not c.ok:
                    print(f"\n[round {r.round_num}] {c.name}\n{c.detail}")

    print(f"\nTotal: {total_passed}/{total_cases} passed across {args.rounds} round(s)")
    print(f"       {total_failed} failure(s)")
    return 0 if total_failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
