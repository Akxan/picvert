"""Real document conversion across docx / xlsx / csv.

Conversion matrix (all directions implemented):

           target →  DOCX   XLSX   CSV
    ┌────────────┬──────┬──────┬──────┐
    │ DOCX       │  ✓¹  │  ✓   │  ✓   │
    │ XLSX       │  ✓   │  ✓¹  │  ✓²  │
    │ CSV        │  ✓   │  ✓   │  ✓¹  │
    └────────────┴──────┴──────┴──────┘
    ¹ same-format ⇒ copy
    ² xlsx → csv writes one CSV per sheet (returns N)
"""
from __future__ import annotations

import csv
import logging
import shutil
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocxDocument
from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Reading helpers — produce a uniform "rows" representation we can write back.
# ---------------------------------------------------------------------------

def _read_xlsx(path: Path) -> dict[str, list[list[str]]]:
    """Return {sheet_name: rows} for every sheet in the workbook.

    Raises ValueError if the workbook contains no readable sheets so callers
    get a clean error instead of a downstream KeyError.
    """
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    sheets: dict[str, list[list[str]]] = {}
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            sheets[name] = [
                ["" if cell is None else str(cell) for cell in row]
                for row in ws.iter_rows(values_only=True)
            ]
    finally:
        wb.close()
    if not sheets:
        raise ValueError(f"xlsx file has no sheets: {path}")
    return sheets


def _read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [list(row) for row in csv.reader(f)]


def _read_docx(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    """Return (paragraphs, tables)."""
    doc: DocxDocument = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = [
        [[cell.text for cell in row.cells] for row in tbl.rows]
        for tbl in doc.tables
    ]
    return paragraphs, tables


# ---------------------------------------------------------------------------
# Writing helpers
# ---------------------------------------------------------------------------

def _write_csv(rows: Iterable[Iterable[str]], path: Path) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)


def _write_xlsx(sheets: dict[str, list[list[str]]], path: Path) -> None:
    wb = Workbook()
    default_ws = wb.active
    first = True
    for name, rows in sheets.items():
        ws = default_ws if first else wb.create_sheet()
        ws.title = (name or "Sheet1")[:31]  # Excel limits sheet name to 31 chars
        for row in rows:
            ws.append(row)
        first = False
    if first:
        # No sheets at all: drop a blank "Sheet1".
        default_ws.title = "Sheet1"
    wb.save(str(path))


def _add_table_to_doc(doc: DocxDocument, rows: list[list[str]]) -> None:
    """Append a Word table built from `rows`. Empty rows ⇒ "(empty)" placeholder."""
    if not rows:
        doc.add_paragraph("(empty)")
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            table.cell(i, j).text = row[j] if j < len(row) else ""


def _write_docx_from_table(rows: list[list[str]], path: Path, title: str | None = None) -> None:
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    _add_table_to_doc(doc, rows)
    doc.save(str(path))


def _write_docx_from_sheets(sheets: dict[str, list[list[str]]], path: Path) -> None:
    doc = Document()
    for name, rows in sheets.items():
        doc.add_heading(name, level=1)
        _add_table_to_doc(doc, rows)
        doc.add_page_break()
    doc.save(str(path))


# ---------------------------------------------------------------------------
# Public dispatcher
# ---------------------------------------------------------------------------

def _is_same_file(a: Path, b: Path) -> bool:
    """True when `a` and `b` point at the same on-disk location."""
    try:
        return a.resolve() == b.resolve()
    except OSError:
        return False


def convert_document_file(
    file_path: Path, output_folder: Path, output_format: str, ext_out: str
) -> int:
    """Convert one document file. Returns the number of files written."""
    src_ext = file_path.suffix.lower()
    out_path = output_folder / (file_path.stem + ext_out)

    if _is_same_file(file_path, out_path):
        logger.info("Skipping %s: source and destination are the same path", file_path)
        return 0

    # Same format → copy.
    if (src_ext, output_format) in {(".docx", "DOCX"), (".xlsx", "XLSX"), (".csv", "CSV")}:
        shutil.copyfile(str(file_path), str(out_path))
        logger.info("Document copied %s → %s", file_path, out_path)
        return 1

    if src_ext == ".csv":
        rows = _read_csv(file_path)
        if output_format == "XLSX":
            _write_xlsx({"Sheet1": rows}, out_path)
        elif output_format == "DOCX":
            _write_docx_from_table(rows, out_path, title=file_path.stem)
        else:
            raise ValueError(f"Unsupported csv → {output_format}")
        logger.info("Document converted %s → %s", file_path, out_path)
        return 1

    if src_ext == ".xlsx":
        sheets = _read_xlsx(file_path)
        if output_format == "CSV":
            # One CSV per sheet — preserves data instead of silently dropping it.
            return _xlsx_to_csv(sheets, file_path, output_folder, ext_out)
        if output_format == "DOCX":
            _write_docx_from_sheets(sheets, out_path)
            logger.info("Document converted %s → %s", file_path, out_path)
            return 1
        raise ValueError(f"Unsupported xlsx → {output_format}")

    if src_ext == ".docx":
        paragraphs, tables = _read_docx(file_path)
        if output_format == "XLSX":
            sheets: dict[str, list[list[str]]] = {}
            for i, t in enumerate(tables, 1):
                sheets[f"Table{i}"] = t
            if paragraphs:
                sheets["Paragraphs"] = [[p] for p in paragraphs]
            if not sheets:
                sheets["Sheet1"] = [[""]]
            _write_xlsx(sheets, out_path)
        elif output_format == "CSV":
            if tables:
                _write_csv(tables[0], out_path)
            else:
                _write_csv([[p] for p in paragraphs], out_path)
        else:
            raise ValueError(f"Unsupported docx → {output_format}")
        logger.info("Document converted %s → %s", file_path, out_path)
        return 1

    raise ValueError(f"Unknown document type: {src_ext}")


def _xlsx_to_csv(
    sheets: dict[str, list[list[str]]],
    file_path: Path,
    output_folder: Path,
    ext_out: str,
) -> int:
    """Write one CSV per sheet.

    Naming is always `<stem>__<sheet>.csv` — including for the single-sheet
    case. Previously single-sheet skipped the suffix, but that produced
    different filenames depending on workbook structure, which collided
    badly when a batch contained both shapes (e.g. a `report.xlsx` with one
    sheet and a sibling `report.csv` would clobber each other).

    Sheet names are sanitised to filesystem-safe characters; collisions
    after sanitisation (e.g. `A/B` and `A_B` both → `A_B`) get a numeric
    suffix so no write is silently dropped.
    """
    used: set[str] = set()
    written = 0
    for sheet_name, rows in sheets.items():
        safe = "".join(c if c.isalnum() or c in "-_." else "_" for c in sheet_name) or "sheet"
        candidate = safe
        n = 1
        while candidate in used:
            n += 1
            candidate = f"{safe}_{n}"
        used.add(candidate)
        out = output_folder / f"{file_path.stem}__{candidate}{ext_out}"
        _write_csv(rows, out)
        logger.info("Document converted %s sheet %r → %s", file_path, sheet_name, out)
        written += 1
    return written
